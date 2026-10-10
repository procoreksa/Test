"""Week 1 content for Pro Core Media (28 posts) + Excel dashboard and Markdown export.

Usage: python3 week1.py  -> writes ProCoreMedia-Content-Dashboard.xlsx and week1.md next to this file.
All posts start as: production=Not started, review=Pending approval, publishing=Not scheduled.
"""
import os
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
START = date(2026, 10, 13)  # proposed launch day (Tuesday) — after approval
SLOTS = {1: "09:00", 2: "13:30", 3: "17:30", 4: "21:00"}  # Asia/Riyadh, refine with Metricool data
WA = "📩 راسلنا على واتساب: [رابط واتساب]"

HT_CORE = "#ProCoreMedia #تسويق_رقمي #السعودية #المنطقة_الشرقية"
HT = {
    "brand": HT_CORE + " #الخبر #الدمام #وكالة_تسويق",
    "edu": HT_CORE + " #نصائح_تسويقية #سوشيال_ميديا #صناعة_المحتوى",
    "ads": HT_CORE + " #إعلانات_ممولة #MetaAds #تسويق_إلكتروني",
    "ai": HT_CORE + " #ذكاء_اصطناعي #AI #تصميم",
    "port": HT_CORE + " #تصميم_إعلانات #هوية_بصرية #إبداع",
}

STYLE = ("ProCore Media primary identity: deep black background, neon gradient accents "
         "(#14C8FF → #2F6BFF → #8B1FE8 → #F2349E → #FF8A00), Poppins/Tajawal typography, generous negative space, "
         "small colour P-mark logo bottom corner.")
GOLD = ("ProCore Media premium identity: black background, metallic gold (#D4AF37 gradient), luxury minimal layout, "
        "small gold P-mark logo.")

# (day, slot, platform, objective, category, format, idea, on_design, caption, cta, hashtags_key, direction, prompt, size)
P = [
 # ---------------- Day 1 — launch ----------------
 (1,1,"FB+IG","Brand Awareness","Brand Awareness","Image (Pinned #1)",
  "إعلان انطلاق الوكالة",
  "ProCore Media\nحيث يلتقي الإبداع بالابتكار",
  "مرحبًا بكم في ProCore Media 🖤 وكالة إعلانات وتسويق رقمي سعودية، نجمع بين الإبداع والذكاء الاصطناعي والبيانات لنصنع لعلامتك حضورًا يليق بها ونتائج تُقاس.\nتابعونا… القادم أجمل.",
  "تابعنا وأرسل لنا رسالة","brand",
  STYLE+" Centered large colour P-mark with soft neon glow, wordmark below, Arabic tagline.",
  "Ultra-clean launch poster, black background, glowing multicolour ribbon letter P with a play-button cut-out at centre, soft volumetric light, premium tech agency aesthetic, lots of negative space, 4:5 — NO text (text added in design).",
  "1080×1350"),
 (1,2,"FB+IG","Agency Services","Agency Services","Carousel 5 (Pinned #2)",
  "ماذا نقدم: 4 خدمات",
  "شريحة 1: 4 خدمات… هدف واحد: نمو علامتك\n2: إدارة السوشيال والتصميم\n3: إعلانات Meta\n4: Reels وفيديو وAI\n5: الهوية البصرية",
  "4 خدمات، هدف واحد: نمو علامتك 🚀\n1) إدارة حسابات التواصل والتصميم\n2) إدارة إعلانات Meta مع تقارير واضحة\n3) Reels وفيديوهات ومحتوى بالذكاء الاصطناعي\n4) تصميم الهوية البصرية\nأي خدمة تحتاجها أولًا؟ اكتبها في التعليقات 👇",
  "اكتب (خدمات) في الرسائل","brand",
  STYLE+" One service per slide: big line icon in neon gradient, service name, 2-line benefit. Consistent grid.",
  "Set of 4 minimal neon line icons on black (grid of squares, megaphone with target, play button with sparkle, pen nib with colour swatches), consistent stroke, gradient cyan-purple-orange — for carousel slides.",
  "1080×1350 ×5"),
 (1,3,"FB+IG","Brand Awareness","AI Marketing Content","Reel 12s",
  "Reel تعريفي: الشعار يتكوّن من الشرائط الملونة",
  "ProCore Media\nCreative · AI · Performance",
  "هكذا تبدأ القصة ✨ ProCore Media… حيث يلتقي الإبداع بالابتكار.",
  "تابع الحساب","ai",
  "Motion graphics: ribbons fly in and assemble the P-mark, light sweep, wordmark reveal, tagline; 12s, music bed (licensed).",
  "Cinematic 3D animation: multicolour silk ribbons (cyan, blue, purple, pink, orange) swirl in darkness and fold together into a letter P with a play-button hole, glossy light sweep, slow push-in, black background, 9:16.",
  "1080×1920"),
 (1,4,"FB+IG","Lead Generation","Lead Generation","Image (Pinned #3)",
  "كيف نبدأ معك: 3 خطوات",
  "3 خطوات وتبدأ\nاستشارة ← خطة ← تنفيذ",
  "البداية سهلة:\n1) مكالمة قصيرة نفهم فيها نشاطك وأهدافك\n2) خطة محتوى وإعلانات مخصصة\n3) تنفيذ وتقرير شهري واضح\nاحجز استشارتك الآن 📩",
  "احجز استشارة عبر واتساب","brand",
  STYLE+" Three numbered steps connected by a gradient line, icons: phone, plan document, rocket.",
  "Minimal 3-step infographic background on black: three glowing nodes connected by a neon gradient line, subtle grid, 4:5, no text.",
  "1080×1350"),
 # ---------------- Day 2 ----------------
 (2,1,"FB+IG","Educational","Educational Content","Carousel 6",
  "5 أخطاء تقتل تفاعل حسابك",
  "5 أخطاء تقتل تفاعل حسابك\n1 نشر بلا هدف · 2 تصاميم مزدحمة · 3 بلا Hook · 4 تجاهل التعليقات · 5 لا توجد دعوة لإجراء",
  "هل تنشر باستمرار والتفاعل ضعيف؟ غالبًا السبب واحد من هذه الأخطاء الخمسة 👇\nاحفظ المنشور وراجع حسابك عليها.",
  "احفظ المنشور + شاركه","edu",
  STYLE+" Slide per mistake: big number, red-ish X icon in gradient, one-line fix in white.",
  "Abstract background: a phone screen showing a social feed with fading heart icons, dark moody lighting, neon edge light, 4:5, no readable text.",
  "1080×1350 ×6"),
 (2,2,"FB+IG","Agency Services","Agency Services","Image",
  "خدمة إدارة السوشيال: ماذا تشمل",
  "إدارة حساباتك علينا\nخطة · تصميم · نشر · تفاعل · تقرير",
  "خلّ حساباتك تشتغل وأنت مركّز على نشاطك 💼\nفي خدمة إدارة السوشيال نجهّز لك: خطة شهرية، تصاميم احترافية، نشر منتظم، متابعة التعليقات والرسائل، وتقرير أداء.",
  "اسأل عن الباقات","brand",
  STYLE+" Phone mockup with a grid of neat branded posts, five small checklist items.",
  "Premium smartphone mockup floating on black, screen shows a tidy colourful Instagram grid (abstract shapes), neon rim light, 4:5.",
  "1080×1350"),
 (2,3,"FB+IG","AI Marketing","AI Marketing Content","Reel 15s",
  "قبل/بعد: صورة منتج عادية ← إعلان بالذكاء الاصطناعي (عينة AI)",
  "قبل ← بعد\nعينة AI",
  "نفس المنتج… نتيجة مختلفة تمامًا 🤯 صورة عادية بالجوال تتحول لإعلان احترافي بمساعدة الذكاء الاصطناعي.\n(عينة توضيحية من إنتاجنا بالذكاء الاصطناعي — ليست لعميل)",
  "تبي نجرّب على منتجك؟ راسلنا","ai",
  "Split-screen wipe transition from plain phone photo to styled AI ad scene; label 'AI Sample' on screen.",
  "Product photography transformation: a simple coffee cup on an office desk transforms into a luxurious advertising scene with dramatic lighting, floating coffee beans and steam, cinematic, 9:16.",
  "1080×1920"),
 (2,4,"FB+IG","Engagement","Interactive Content","Image (Question)",
  "سؤال: وش أكبر تحدي يواجهك في التسويق؟",
  "وش أكبر تحدي يواجهك في التسويق؟\nA أفكار المحتوى · B التصميم · C الإعلانات · D الوقت",
  "سؤالنا الليلة 👇 اختر حرفًا واحدًا في التعليقات، وبنجاوب على أكثر تحدٍّ تكرر في منشور قادم.",
  "علّق بحرف","brand",
  STYLE+" Four answer tiles A–D in gradient outline boxes.",
  "Abstract 3D question mark made of neon ribbons on black background, soft glow, 4:5.",
  "1080×1350"),
 # ---------------- Day 3 ----------------
 (3,1,"FB+IG","Educational","Digital Marketing Tips","Carousel 4",
  "الفرق بين Reach وImpressions",
  "Reach vs Impressions\nكم شخص شافك؟ وكم مرة؟",
  "Reach = عدد الأشخاص الذين شاهدوا منشورك.\nImpressions = عدد مرات الظهور (الشخص الواحد قد يشاهده أكثر من مرة).\nلماذا يهمك الفرق؟ لأنه يكشف هل تصل لجمهور جديد أم تكرر الظهور لنفس الأشخاص.",
  "احفظ المنشور","edu",
  STYLE+" Simple diagram: people icons vs eye icons with counters.",
  "Clean infographic illustration: people silhouettes on the left, repeated eye icons on the right, neon gradient lines, black background, 4:5, no text.",
  "1080×1350 ×4"),
 (3,2,"FB+IG","Agency Services","Agency Services","Image",
  "خدمة إعلانات Meta",
  "إعلاناتك على Meta\nاستهداف · إبداع · قياس",
  "الإعلان الناجح ليس صدفة 🎯 نبني حملاتك على Facebook وInstagram خطوة بخطوة: جمهور واضح، إعلانات مصممة للتحويل، متابعة يومية، وتقرير يوضح أين صُرفت كل ريال.",
  "اطلب خطة إعلانية","ads",
  STYLE+" Dashboard-style visual with abstract bars and a target icon.",
  "Futuristic ads dashboard floating in darkness, abstract bar charts and target rings glowing in cyan-purple-orange gradients, no numbers, 4:5.",
  "1080×1350"),
 (3,3,"FB+IG","AI Marketing","AI Marketing Content","Reel 20s",
  "كيف نستخدم الذكاء الاصطناعي في المحتوى (3 استخدامات)",
  "3 طرق نستخدم فيها AI\nأفكار · صور · فيديو",
  "الذكاء الاصطناعي ما يلغي الإبداع… يسرّعه ⚡ نستخدمه في: 1) توليد أفكار أكثر 2) إنتاج صور إعلانية 3) مشاهد فيديو — والقرار النهائي دائمًا لفريقنا.",
  "تابعنا لمزيد من نصائح AI","ai",
  "Fast-cut motion graphics with three chapters, kinetic Arabic typography, screen-style UI mockups (no real brand UIs).",
  "Abstract futuristic creative studio: holographic panels showing image thumbnails and video timelines, neon gradient light, camera slowly pushing in, 9:16.",
  "1080×1920"),
 (3,4,"FB+IG","Portfolio","Portfolio & Creative Showcase","Carousel 3 (Concept)",
  "Concept: حملة لمقهى مختص",
  "Concept Project\nحملة افتتاح مقهى",
  "مشروع تصوري (Concept) ☕ تخيّلنا حملة افتتاح لمقهى مختص في الخبر: فكرة، تصاميم، ورسالة واحدة واضحة. هذا نموذج إبداعي يوضح أسلوبنا — وليس عملًا لعميل.",
  "تبي حملة لنشاطك؟ راسلنا","port",
  STYLE+" Label 'Concept Project' visible on every slide.",
  "Moody specialty coffee shop advertising visual: barista pouring latte art under warm spotlight, dark background, cinematic shallow depth of field, 4:5.",
  "1080×1350 ×3"),
 # ---------------- Day 4 ----------------
 (4,1,"FB+IG","Educational","Digital Marketing Tips","Image",
  "أفضل وقت للنشر؟ الجواب في بياناتك",
  "أفضل وقت للنشر؟\nبياناتك تعرف الجواب",
  "ما في وقت سحري يناسب الجميع ⏰ ابدأ بتجربة أوقات مختلفة لمدة أسبوعين، ثم افتح Insights وشوف متى جمهورك متصل أكثر. نحن نعتمد هذا الأسلوب مع كل حساب.",
  "جرّبها وأخبرنا","edu",
  STYLE+" Clock icon merging with a bar chart.",
  "Minimal neon clock whose hands turn into rising bar charts, black background, 4:5.",
  "1080×1350"),
 (4,2,"FB+IG","Agency Services","Agency Services","Image",
  "خدمة Reels والفيديو",
  "Reels تخطف الانتباه\nفكرة · تصوير/AI · مونتاج",
  "الفيديو القصير هو أقوى أداة وصول اليوم 🎬 ننتج لك Reels من الفكرة إلى المونتاج — بالتصوير أو بمشاهد مولّدة بالذكاء الاصطناعي.",
  "اطلب فكرة Reel لنشاطك","ai",
  STYLE+" Vertical phone frame with film-strip and play icon.",
  "Vertical phone with a glowing play button and film strip flowing out of the screen, neon gradient light trails, black background, 4:5.",
  "1080×1350"),
 (4,3,"FB+IG","AI Marketing","AI Creative Showcase","Reel 15s (AI Sample)",
  "مشهد سينمائي AI لمطعم (عينة)",
  "عينة AI",
  "مشهد سينمائي أُنتج بالكامل بالذكاء الاصطناعي 🍽️ هذا النوع من المشاهد يمكن استخدامه في إعلانات المطاعم بتكلفة أقل من التصوير التقليدي.\n(عينة توضيحية — ليست لعميل)",
  "تبي مشهد لنشاطك؟ راسلنا","ai",
  "Disclose 'AI-generated sample' on screen and enable the platform AI label.",
  "Cinematic slow-motion shot of a sizzling grilled steak on a cast-iron plate, flames and smoke, dark restaurant ambience, warm rim light, slow push-in, 9:16.",
  "1080×1920"),
 (4,4,"FB+IG","Engagement","Interactive Content","Image (A/B)",
  "أي تصميم أفضل: A أو B؟",
  "A أو B؟",
  "تصميمان لنفس الإعلان… أيهما يجذبك أكثر؟ 👀 علّق بـ A أو B — وبنكشف بكرة ليش الفرق مهم في الإعلانات.",
  "علّق A أو B","port",
  STYLE+" Two ad variants side by side (one cluttered, one clean) — both labeled Concept.",
  "Two side-by-side abstract ad layouts on black: left busy and cluttered, right clean and minimal with neon accent, 4:5.",
  "1080×1350"),
 # ---------------- Day 5 ----------------
 (5,1,"FB+IG","Educational","Educational Content","Carousel 5",
  "كيف تكتب Hook يوقف التمرير",
  "أول 3 ثوانٍ تقرر كل شيء\n5 صيغ Hook مجرّبة",
  "أول سطر في منشورك هو الإعلان الحقيقي ✍️ 5 صيغ تساعدك: سؤال مباشر، رقم مفاجئ، خطأ شائع، وعد واضح، قصة قصيرة. أي صيغة تستخدمها أكثر؟",
  "احفظ + جرّب","edu",
  STYLE+" One formula per slide with an Arabic example.",
  "Thumb stopping a scrolling feed on a phone, motion blur on the feed, sharp highlight on one post, neon light, 4:5.",
  "1080×1350 ×5"),
 (5,2,"FB+IG","Agency Services","Agency Services","Image",
  "خدمة الهوية البصرية",
  "هويتك أول انطباع\nشعار · ألوان · خطوط · تطبيقات",
  "قبل ما يقرأ العميل كلمة واحدة، هويتك تكون تكلمت عنك 🎨 نصمم لك شعارًا ونظام ألوان وخطوط وتطبيقات جاهزة للسوشيال والمطبوعات.",
  "اطلب عرض هوية","port",
  STYLE+" Brand board layout: logo placeholder, colour chips, type specimen.",
  "Elegant brand identity board mockup on black: stationery, business cards, colour swatches, abstract logo placeholder, soft studio light, 4:5.",
  "1080×1350"),
 (5,3,"FB+IG","Portfolio","Behind the Scenes","Reel 20s",
  "خلف الكواليس: كيف نبني منشورًا",
  "من الفكرة إلى النشر",
  "من فكرة على ورقة إلى منشور جاهز 🛠️ هكذا نشتغل: بحث، فكرة، تصميم، مراجعة، ثم نشر في الوقت المناسب.",
  "تابعنا","brand",
  "Screen-recording style montage of our own design process (internal work), sped up, with step labels.",
  "Top-down shot of a designer's desk at night: sketchbook, tablet with design software, coffee, neon desk lamp, time-lapse feel, 9:16.",
  "1080×1920"),
 (5,4,"FB+IG","Lead Generation","Lead Generation","Image ⚠️",
  "عرض: مراجعة مجانية لحسابك (يحتاج موافقة)",
  "مراجعة مجانية لحسابك\n+ 3 توصيات عملية",
  "نقدم مراجعة مجانية لحسابات نشاطك على Instagram وFacebook مع 3 توصيات عملية تقدر تطبقها فورًا ✅ العرض لأول 10 منشآت.",
  "أرسل رابط حسابك على واتساب","brand",
  STYLE+" Magnifying glass over a phone profile grid.",
  "Glowing magnifying glass hovering over a smartphone showing a social media profile grid, neon gradient light, black background, 4:5.",
  "1080×1350"),
 # ---------------- Day 6 ----------------
 (6,1,"FB+IG","Brand Awareness","Brand Awareness","Image",
  "3 مبادئ نشتغل بها",
  "وضوح · إبداع · أرقام",
  "3 مبادئ لا نتنازل عنها: وضوح في الخطة، إبداع في التنفيذ، وأرقام تُقاس في النهاية 🖤",
  "تابعنا","brand",
  STYLE+" Three pillars with icons.",
  "Three glowing vertical pillars of light in cyan, purple and orange on a black stage, minimal, 4:5.",
  "1080×1350"),
 (6,2,"FB+IG","Portfolio","Real Estate Portfolio","Image (Concept · Gold)",
  "Concept: إعلان فلل فاخرة",
  "Concept Project\nفلل بإطلالة على البحر",
  "مشروع تصوري (Concept) 🏡 إعلان لمشروع فلل فاخر بأسلوب الهوية الذهبية. نموذج يوضح أسلوبنا في التسويق العقاري — وليس عملًا لعميل.",
  "عندك مشروع عقاري؟ راسلنا","port",
  GOLD+" Label 'Concept Project'.",
  "Luxury modern villa at golden hour with infinity pool overlooking the Arabian Gulf, warm cinematic light, architectural photography style, 4:5.",
  "1080×1350"),
 (6,3,"FB+IG","AI Marketing","AI Creative Showcase","Reel 15s (AI Sample)",
  "قبل/بعد هوية بصرية (عينة)",
  "قبل ← بعد\nConcept",
  "كيف تتغير صورة النشاط لما تتطور هويته ✨ (مشروع تصوري للتوضيح — ليس لعميل)",
  "اطلب عرض هوية","port",
  "Morph from a generic logo/storefront to a refined identity; on-screen 'Concept'.",
  "A plain shopfront at night transforms into a premium branded storefront with elegant signage and warm lighting, smooth morph, 9:16.",
  "1080×1920"),
 (6,4,"FB+IG","Engagement","Interactive Content","Image",
  "اكتب نشاطك ونقترح لك فكرة محتوى",
  "اكتب نوع نشاطك 👇\nونقترح لك فكرة محتوى",
  "تحدٍّ اليوم 💡 اكتب نوع نشاطك في التعليقات (مطعم، عيادة، عقار، متجر…) وفريقنا يقترح عليك فكرة منشور مجانًا.",
  "علّق بنوع نشاطك","brand",
  STYLE+" Speech bubbles with business icons.",
  "Floating neon speech bubbles containing simple icons (fork and knife, house, tooth, shopping bag) on black, 4:5.",
  "1080×1350"),
 # ---------------- Day 7 ----------------
 (7,1,"FB+IG","Educational","Digital Marketing Tips","Carousel 4",
  "ميزانية إعلانات Meta: من وين تبدأ؟",
  "ميزانية إعلاناتك\nمن وين تبدأ؟",
  "سؤال نسمعه كثيرًا 💰 ابدأ بميزانية اختبار صغيرة، جرّب أكثر من إعلان، ثم ضاعف الصرف على الأفضل أداءً. القرار يكون بالأرقام لا بالتخمين.",
  "احفظ المنشور","ads",
  STYLE+" Steps: test → learn → scale.",
  "Coins turning into upward glowing arrows on a dark background, neon gradient, 4:5.",
  "1080×1350 ×4"),
 (7,2,"FB+IG","Promotional","Promotional Content","Image ⚠️",
  "باقاتنا: Start · Growth · Pro (بعد اعتماد الأسعار)",
  "باقات تناسب كل مرحلة\nStart · Growth · Pro",
  "سواء كنت تبدأ أو تبي تكبر بسرعة، عندنا باقة تناسبك 📦 اطلب تفاصيل الباقات وقارن بنفسك.",
  "اطلب الباقات","brand",
  STYLE+" Three package cards, middle highlighted.",
  "Three floating glass cards on black with neon edges, the middle one taller and brighter, 4:5.",
  "1080×1350"),
 (7,3,"FB+IG","Portfolio","Sports & Fitness Portfolio","Reel 15s (AI Sample)",
  "Concept: إعلان أكاديمية سباحة للأطفال",
  "Concept · AI Sample",
  "تخيّلنا إعلانًا لأكاديمية سباحة للأطفال 🏊 مشاهد مولّدة بالذكاء الاصطناعي توضح كيف يمكن عرض تجربة الأطفال بشكل جذاب.\n(نموذج تصوري — ليس لعميل)",
  "عندك أكاديمية أو نادٍ؟ راسلنا","ai",
  "Avoid identifiable real children; stylised/back-view swimmers only; on-screen 'Concept · AI'.",
  "Bright indoor swimming pool, underwater cinematic shot of a child swimmer from behind with goggles, sun rays through water, joyful mood, 9:16.",
  "1080×1920"),
 (7,4,"FB+IG","Lead Generation","Lead Generation","Image",
  "خلاصة الأسبوع + احجز استشارة",
  "أسبوعنا الأول ✨\nجاهزين لنشتغل على علامتك",
  "أسبوع أول مليء بالأفكار 🙌 شكرًا لكل من تابعنا وتفاعل. إذا تبي نخطط محتوى وإعلانات نشاطك للشهر القادم، احجز استشارتك الآن.",
  "احجز استشارة","brand",
  STYLE+" Collage of the week's visuals in a grid.",
  "Mosaic of colourful abstract creative tiles glowing on black, 4:5.",
  "1080×1350"),
]

HEADERS = ["Content ID", "Day", "Post #", "Date", "Publishing Time", "Platform", "Objective", "Content Category",
           "Content Format", "Idea", "Text on Design", "Caption (Instagram)", "Caption (Facebook)", "Call to Action",
           "Hashtags", "Creative Direction", "AI Prompt", "Size", "Creative File", "Production Status",
           "Approval Status", "Publishing Status", "Post URL", "Reach", "Impressions", "Engagement", "Comments",
           "Shares", "Saves", "Profile Visits", "Leads Generated", "Notes"]


def rows():
    for (day, slot, plat, obj, cat, fmt, idea, ondesign, cap, cta, htk, direction, prompt, size) in P:
        d = START + timedelta(days=day - 1)
        cid = f"PCM-W1-D{day}-P{slot}"
        needs = "⚠️" in fmt or "⚠️" in idea
        ig = f"{cap}\n\n👉 {cta}\n.\n{HT[htk]}"
        fb = f"{cap}\n\n👉 {cta}\n{WA}"
        yield [cid, day, slot, d.isoformat(), SLOTS[slot], plat, obj, cat, fmt, idea, ondesign, ig, fb, cta,
               HT[htk], direction, prompt, size, "", "Not started",
               "Pending approval" + (" — needs offer/pricing approval" if needs else ""), "Not scheduled",
               "", "", "", "", "", "", "", "", "", ""]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Content Week 1"
    ws.sheet_view.rightToLeft = False
    ws.append(HEADERS)
    for r in rows():
        ws.append(r)
    head = PatternFill("solid", fgColor="111111")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    widths = {"A": 18, "D": 12, "E": 10, "J": 30, "K": 30, "L": 50, "M": 50, "P": 40, "Q": 50}
    for col in ws.iter_cols(min_row=1, max_row=1):
        ws.column_dimensions[col[0].column_letter].width = widths.get(col[0].column_letter, 16)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    for col, opts in (("T", "Not started,In production,Ready"), ("U", "Pending approval,Approved,Changes requested,Rejected"),
                      ("V", "Not scheduled,Scheduled,Published,Failed,Paused")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}500")

    k = wb.create_sheet("KPI 30 days")
    k.append(["KPI", "Target (estimate)", "Actual", "Source"])
    for row in [("Posts published successfully", "120", "", "Metricool"), ("Schedule adherence", "≥ 95%", "", "Metricool"),
                ("Follower growth (IG+FB)", "300–600", "", "Metricool"), ("Reach", "15k–40k", "", "Metricool"),
                ("Engagement rate", "3–6%", "", "Metricool"), ("Profile visits", "1,000+", "", "Instagram Insights"),
                ("WhatsApp clicks", "60+", "", "Meta Insights"), ("Qualified leads", "10–20", "", "Manual log"),
                ("Client meetings", "4–8", "", "Manual log"), ("Proposals sent", "3–6", "", "Manual log"),
                ("Clients acquired", "1–2", "", "Manual log"), ("Cost per lead (if ads run)", "TBD after test", "", "Meta Ads")]:
        k.append(list(row))
    for c in k[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
    for col, w in (("A", 32), ("B", 18), ("C", 14), ("D", 20)):
        k.column_dimensions[col].width = w
    k.append([])
    k.append(["Targets are estimates for a brand-new account and are not results."])
    wb.save(os.path.join(HERE, "ProCoreMedia-Content-Dashboard.xlsx"))

    with open(os.path.join(HERE, "week1.md"), "w") as f:
        f.write("# ProCore Media — محتوى الأسبوع الأول (28 منشورًا)\n\n"
                f"تاريخ البدء المقترح: {START.isoformat()} · الأوقات بتوقيت الرياض: 09:00 · 13:30 · 17:30 · 21:00\n\n"
                "الحالة الافتراضية لكل منشور: الإنتاج = لم يبدأ · المراجعة = بانتظار الموافقة · النشر = غير مجدول\n\n")
        for r in rows():
            f.write(f"## {r[0]} — {r[3]} {r[4]} — {r[8]}\n\n"
                    f"- **الهدف:** {r[6]} · **التصنيف:** {r[7]} · **المنصة:** {r[5]} · **المقاس:** {r[17]}\n"
                    f"- **الفكرة:** {r[9]}\n- **النص داخل التصميم:** {r[10].replace(chr(10), ' / ')}\n"
                    f"- **CTA:** {r[13]}\n- **Creative Direction:** {r[15]}\n- **AI Prompt:** `{r[16]}`\n"
                    f"- **الحالة:** {r[19]} · {r[20]} · {r[21]}\n\n"
                    f"**Caption — Instagram:**\n\n```\n{r[11]}\n```\n\n**Caption — Facebook:**\n\n```\n{r[12]}\n```\n\n")
    print("ok", len(P), "posts")


if __name__ == "__main__":
    main()
